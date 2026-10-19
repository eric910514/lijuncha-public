"""把作者下載的官方姓名資料轉成 config/names/names_foreign.jsonl（格式見〈原文姓名資料檔規格〉）。

用法（在程式資料夾）：
    python -m pip install openpyxl
    python tools/make_names_foreign.py --raw <NAMES 資料夾>

<NAMES 資料夾> 就是作者整理的 NAMES（裡面有 FRANCE、UK、USA、CANADA、AU…）。目前會轉：

| 國家 | 名字 | 姓氏 |
| 法國 | INSEE Fichier des prénoms（出生年 1955–2007 合計） | INSEE 姓氏檔（FR_lastnames.csv） |
| 英國 | 蘇格蘭 2025 年新生兒名字 | 蘇格蘭 2025 年最常見姓氏 |
| 美國 | SSA 嬰兒名字（出生年 1955–2007 合計） | 沒有（Census 姓氏 CSV 還沒下載） |
| 加拿大 | 省份名字檔（出生年 1955–2007 合計） | 沒有 |
| 澳大利亞 | NSW 2025 年前 100 名 | 沒有 |

每國每性別取人數最多的前 100 個名字，姓氏取前 100 個；之後 gen 會再套 names_min_count（人數少於 100 不用）。
缺姓氏的國家，gen 會列為「缺姓名」，不會拿來產生正式人物。轉完會印出每國每性別的筆數，請人工看過一次。"""
import argparse
import csv
import io
import json
import os
import zipfile

TOP = 100
YEARS = (1955, 2007)        # 人物年齡約 18 到 70 歲：出生年 1955 到 2007


def nice(s):
    """全大寫或全小寫的名字轉成一般寫法：JEAN-PIERRE → Jean-Pierre，le gall → Le Gall。"""
    s = s.strip()
    return s if (not s.isupper() and not s.islower()) else s.title()


def src(sid, citation, url, locator, license_):
    return {"source_id": sid, "citation": citation, "url": url, "locator": locator, "license": license_}


def top(counts, n=TOP):
    return sorted(counts.items(), key=lambda kv: -kv[1])[:n]


def france(raw, out):
    z = zipfile.ZipFile(os.path.join(raw, "FRANCE", "FR_prenoms-2025-nat_csv.zip"))
    name = [n for n in z.namelist() if n.endswith(".csv")][0]
    counts = {"male": {}, "female": {}}
    for r in csv.DictReader(io.TextIOWrapper(z.open(name), "utf-8"), delimiter=";"):
        if not r["periode"].isdigit() or not YEARS[0] <= int(r["periode"]) <= YEARS[1] or r["prenom"].startswith("_"):
            continue
        lean = {"1": "male", "2": "female"}[r["sexe"]]
        counts[lean][nice(r["prenom"])] = counts[lean].get(nice(r["prenom"]), 0) + int(r["valeur"])
    s = src("INSEE-PRENOMS-2025", "INSEE, Fichier des prénoms (édition 2025), France entière", "https://www.insee.fr/",
            f"prenoms-2025-nat.csv，出生年 {YEARS[0]}–{YEARS[1]} 合計", "Licence Ouverte 2.0")
    for lean, c in counts.items():
        out += [{"record_type": "given_name", "name": n, "origin_country": "法國", "gender_lean": lean, "count": v, "source": s}
                for n, v in top(c)]
    sur = {}
    with open(os.path.join(raw, "FRANCE", "FR_lastnames.csv"), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            sur[nice(r["lastname"])] = sur.get(nice(r["lastname"]), 0) + int(r["occurrences"])
    s = src("INSEE-NOMS", "INSEE, noms de famille (fichier des décès 1970–2024)", "https://www.insee.fr/",
            "FR_lastnames.csv", "Licence Ouverte 2.0")
    out += [{"record_type": "surname", "name": n, "origin_country": "法國", "gender_lean": None, "count": v, "source": s}
            for n, v in top(sur)]
    out.append({"record_type": "naming_rule", "name": None, "origin_country": "法國", "order": "given_first", "separator": " ",
                "source": s})


def _sheet_rows(path, sheet):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True)[sheet]
    return list(ws.iter_rows(values_only=True))


def uk(raw, out):
    path = os.path.join(raw, "UK", "SCOTLAND_all-names-given-to-babies-in-2025.xlsx")
    s = src("NRS-BABYNAMES-2025", "National Records of Scotland, Babies' first names 2025", "https://www.nrscotland.gov.uk/",
            "all-names-given-to-babies-in-2025.xlsx（以蘇格蘭資料代表英國）", "OGL v3.0")
    for sheet, lean in (("Boys", "male"), ("Girls", "female")):
        c = {}
        for r in _sheet_rows(path, sheet):
            if len(r) >= 3 and isinstance(r[2], (int, float)) and isinstance(r[1], str):
                c[nice(r[1])] = c.get(nice(r[1]), 0) + int(r[2])
        out += [{"record_type": "given_name", "name": n, "origin_country": "英國", "gender_lean": lean, "count": v, "single_year": True, "source": s}
                for n, v in top(c)]
    path = os.path.join(raw, "UK", "SCOTLAND_most-common-surnames-2025.xlsx")
    s = src("NRS-SURNAMES-2025", "National Records of Scotland, Most Common Surnames 2025", "https://www.nrscotland.gov.uk/",
            "All Surnames for 2025（以蘇格蘭資料代表英國）", "OGL v3.0")
    c = {}
    for r in _sheet_rows(path, "All Surnames for 2025"):
        if len(r) >= 4 and isinstance(r[3], (int, float)) and isinstance(r[2], str):
            c[nice(r[2])] = c.get(nice(r[2]), 0) + int(r[3])
    out += [{"record_type": "surname", "name": n, "origin_country": "英國", "gender_lean": None, "count": v, "single_year": True, "source": s}
            for n, v in top(c)]
    out.append({"record_type": "naming_rule", "name": None, "origin_country": "英國", "order": "given_first", "separator": " ",
                "source": s})


def usa(raw, out):
    folder = os.path.join(raw, "USA")
    zpath = [os.path.join(folder, n) for n in os.listdir(folder) if n.endswith(".zip")][0]
    z = zipfile.ZipFile(zpath)
    counts = {"male": {}, "female": {}}
    for y in range(YEARS[0], YEARS[1] + 1):
        for line in z.read(f"yob{y}.txt").decode("utf-8").splitlines():
            n, sx, k = line.split(",")
            lean = "male" if sx == "M" else "female"
            counts[lean][n] = counts[lean].get(n, 0) + int(k)
    s = src("SSA-BABYNAMES", "U.S. Social Security Administration, Baby Names from Social Security Card Applications",
            "https://catalog.data.gov/dataset/baby-names-from-social-security-card-applications-national-data",
            f"yob{YEARS[0]}–yob{YEARS[1]}.txt 合計", "Public Domain")
    for lean, c in counts.items():
        out += [{"record_type": "given_name", "name": n, "origin_country": "美國", "gender_lean": lean, "count": v, "source": s}
                for n, v in top(c)]
    out.append({"record_type": "naming_rule", "name": None, "origin_country": "美國", "order": "given_first", "separator": " ",
                "source": s})


def canada(raw, out):
    s = src("CA-BABYNAMES", "Canadian provincial baby names (province and licence to be recorded)", "",
            f"baby_names_-_female_.csv、baby_names_-_male.csv，出生年 {YEARS[0]}–{YEARS[1]} 合計", "待補記")
    for fn, lean in (("baby_names_-_female_.csv", "female"), ("baby_names_-_male.csv", "male")):
        c = {}
        with open(os.path.join(raw, "CANADA", fn), encoding="utf-8-sig") as f:
            for r in csv.reader(f):
                if r and r[0].isdigit() and YEARS[0] <= int(r[0]) <= YEARS[1]:
                    c[nice(r[1])] = c.get(nice(r[1]), 0) + int(r[2])
        out += [{"record_type": "given_name", "name": n, "origin_country": "加拿大", "gender_lean": lean, "count": v, "source": s}
                for n, v in top(c)]
    out.append({"record_type": "naming_rule", "name": None, "origin_country": "加拿大", "order": "given_first", "separator": " ",
                "source": s})


def australia(raw, out):
    s = src("NSW-BABYNAMES-2025", "NSW Registry of Births, Deaths & Marriages, Popular Baby Names 2025", "https://www.nsw.gov.au/",
            "female-top-100-2025.xlsx、male-top-100-2025.xlsx", "CC BY 4.0")
    for fn, lean in (("female-top-100-2025.xlsx", "female"), ("male-top-100-2025.xlsx", "male")):
        import openpyxl
        ws = openpyxl.load_workbook(os.path.join(raw, "AU", fn), read_only=True).worksheets[0]
        c = {nice(r[0]): int(r[1]) for r in ws.iter_rows(values_only=True) if isinstance(r[1], (int, float)) and r[0]}
        out += [{"record_type": "given_name", "name": n, "origin_country": "澳大利亞", "gender_lean": lean, "count": v, "single_year": True, "source": s}
                for n, v in top(c)]
    out.append({"record_type": "naming_rule", "name": None, "origin_country": "澳大利亞", "order": "given_first", "separator": " ",
                "source": s})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True, help="作者整理的 NAMES 資料夾")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                                  "config", "names", "names_foreign.jsonl"))
    a = ap.parse_args()
    out = []
    for fn in (france, uk, usa, canada, australia):
        try:
            fn(a.raw, out)
        except FileNotFoundError as e:
            print(f"略過 {fn.__name__}：找不到 {e.filename}")
    with open(a.out, "w", encoding="utf-8") as f:
        for i, r in enumerate(out, 1):
            f.write(json.dumps({"id": f"names_foreign-{i:05d}", **r}, ensure_ascii=False) + "\n")
    print(f"寫入 {len(out)} 筆 → {a.out}")
    summary = {}
    for r in out:
        key = (r["origin_country"], r["record_type"], r.get("gender_lean"))
        summary[key] = summary.get(key, [0, 0])
        summary[key][0] += 1
        summary[key][1] += (r.get("count") or 0) >= 100
    for (c, t, g), (n, ok) in sorted(summary.items(), key=str):
        if t != "naming_rule":
            print(f"  {c} {t} {g or ''}：{n} 筆，其中人數 ≥ 100 的 {ok} 筆")


if __name__ == "__main__":
    main()
